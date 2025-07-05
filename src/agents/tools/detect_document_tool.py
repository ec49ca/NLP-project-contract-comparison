"""
Document Detection Tool

This tool detects document types and classifies their structure (structured, semi-structured, unstructured, scanned).
"""

import logging
import mimetypes
from typing import Dict, Any, Tuple
from pathlib import Path
import re
import json
import os

logger = logging.getLogger(__name__)


class DetectDocumentTool:
    """Tool for detecting document types and structure classification"""

    def __init__(self):
        self.name = "detect_document_type"
        self.description = "Detect document type and classify structure"

    async def execute_tool(self, arguments: Dict[str, Any]) -> Any:
        """Execute the document detection tool"""
        try:
            document_content = arguments.get("document_content", "")
            filename = arguments.get("filename", "")
            content_type = arguments.get("content_type", "")
            file_path = arguments.get("file_path", "")

            if not document_content and not file_path:
                return {
                    "error": "Missing document_content or file_path parameter",
                    "success": False,
                }

            # If file_path is provided, try to read and extract content
            extracted_content = ""
            actual_content = document_content

            if file_path and os.path.exists(file_path):
                extracted_content, extraction_success = self._extract_file_content(
                    file_path
                )
                if extraction_success:
                    actual_content = extracted_content
                    logger.info(f"Successfully extracted content from {file_path}")
                else:
                    logger.warning(
                        f"Could not extract content from {file_path}, using filename for detection"
                    )

            # Multi-layer detection approach
            file_type, confidence = self._detect_file_type(
                actual_content, filename, content_type, file_path
            )

            # Enhanced structure classification using actual content
            structure_type = self._classify_structure_with_content(
                file_type, actual_content, filename
            )

            # Scanned document detection
            is_scanned = self._detect_scanned_content(file_type, actual_content)
            if is_scanned:
                structure_type = "scanned"

            # Additional metadata
            metadata = self._generate_metadata(
                actual_content, filename, file_path, content_type, file_type
            )

            # Add content extraction info
            if extracted_content:
                metadata["content_extracted"] = True
                metadata["extracted_content_length"] = len(extracted_content)
            else:
                metadata["content_extracted"] = False

            # Calculate structure rating out of 7
            structure_rating = self._calculate_structure_rating(structure_type)

            result = {
                "document_type": file_type,
                "structure_type": structure_type,
                "structure_rating": structure_rating,
                "confidence": confidence,
                "is_scanned": is_scanned,
                "metadata": metadata,
                "success": True,
            }

            return result

        except Exception as e:
            logger.error(f"Error in detect_document_tool: {e}")
            return self._fallback_response()

    def _extract_file_content(self, file_path: str) -> Tuple[str, bool]:
        """Extract text content from various file types"""
        try:
            file_ext = Path(file_path).suffix.lower()

            if file_ext == ".docx":
                return self._extract_docx_content(file_path)
            elif file_ext in [".txt", ".md", ".html", ".xml", ".json", ".csv"]:
                return self._extract_text_content(file_path)
            else:
                logger.warning(
                    f"Unsupported file type for content extraction: {file_ext}"
                )
                return "", False

        except Exception as e:
            logger.error(f"Error extracting content from {file_path}: {e}")
            return "", False

    def _extract_docx_content(self, file_path: str) -> Tuple[str, bool]:
        """Extract text content from DOCX files"""
        try:
            # Try importing python-docx
            try:
                from docx import Document
            except ImportError:
                logger.warning(
                    "python-docx not available, falling back to basic analysis"
                )
                return "", False

            doc = Document(file_path)

            # Extract all text content
            full_text = []

            # Extract paragraph text
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    full_text.append(paragraph.text.strip())

            # Extract table content
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        full_text.append(" | ".join(row_text))

            content = "\n".join(full_text)
            return content, True

        except Exception as e:
            logger.error(f"Error extracting DOCX content: {e}")
            return "", False

    def _extract_text_content(self, file_path: str) -> Tuple[str, bool]:
        """Extract content from plain text files"""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return content, True
        except Exception as e:
            logger.error(f"Error reading text file: {e}")
            return "", False

    def _detect_file_type(
        self, content: str, filename: str, content_type: str, file_path: str
    ) -> Tuple[str, float]:
        """Multi-layer file type detection"""

        # Layer 1: Extension-based detection
        ext_type, ext_confidence = self._detect_by_extension(filename or file_path)

        # Layer 2: MIME type detection
        mime_type, mime_confidence = self._detect_by_mime_type(
            content_type, filename or file_path
        )

        # Layer 3: Content analysis
        content_type_detected, content_confidence = self._detect_by_content(content)

        # Choose the detection with highest confidence
        detections = [
            (ext_type, ext_confidence),
            (mime_type, mime_confidence),
            (content_type_detected, content_confidence),
        ]

        # Return the detection with highest confidence
        best_detection = max(detections, key=lambda x: x[1])
        return best_detection

    def _detect_by_extension(self, filename: str) -> Tuple[str, float]:
        """Detect file type by extension"""
        if not filename:
            return "unknown", 0.0

        ext = Path(filename).suffix.lower()

        extension_map = {
            # Documents
            ".pdf": ("pdf", 0.9),
            ".docx": ("docx", 0.9),
            ".doc": ("doc", 0.9),
            ".odt": ("odt", 0.9),
            ".rtf": ("rtf", 0.9),
            # Spreadsheets
            ".xlsx": ("xlsx", 0.9),
            ".xls": ("xls", 0.9),
            ".csv": ("csv", 0.9),
            ".ods": ("ods", 0.9),
            # Presentations
            ".pptx": ("pptx", 0.9),
            ".ppt": ("ppt", 0.9),
            ".odp": ("odp", 0.9),
            # Web and markup
            ".html": ("html", 0.9),
            ".htm": ("html", 0.9),
            ".xml": ("xml", 0.9),
            ".json": ("json", 0.9),
            # Text
            ".txt": ("text", 0.8),
            ".md": ("markdown", 0.8),
            ".rst": ("rst", 0.8),
            # Images
            ".jpg": ("image", 0.9),
            ".jpeg": ("image", 0.9),
            ".png": ("image", 0.9),
            ".tiff": ("image", 0.9),
            ".tif": ("image", 0.9),
            ".bmp": ("image", 0.9),
            ".gif": ("image", 0.9),
            ".webp": ("image", 0.9),
            # Archives
            ".zip": ("archive", 0.9),
            ".rar": ("archive", 0.9),
            ".7z": ("archive", 0.9),
        }

        return extension_map.get(ext, ("unknown", 0.0))

    def _detect_by_mime_type(
        self, content_type: str, filename: str
    ) -> Tuple[str, float]:
        """Detect file type using MIME type"""
        try:
            mime_type = content_type

            # If no content_type provided, try to guess from filename
            if not mime_type and filename:
                mime_type, _ = mimetypes.guess_type(filename)

            if not mime_type:
                return "unknown", 0.0

            mime_map = {
                "application/pdf": ("pdf", 0.8),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
                    "docx",
                    0.8,
                ),
                "application/msword": ("doc", 0.8),
                "application/vnd.oasis.opendocument.text": ("odt", 0.8),
                "application/rtf": ("rtf", 0.8),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": (
                    "xlsx",
                    0.8,
                ),
                "application/vnd.ms-excel": ("xls", 0.8),
                "text/csv": ("csv", 0.8),
                "application/vnd.oasis.opendocument.spreadsheet": ("ods", 0.8),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation": (
                    "pptx",
                    0.8,
                ),
                "application/vnd.ms-powerpoint": ("ppt", 0.8),
                "text/html": ("html", 0.8),
                "application/xml": ("xml", 0.8),
                "text/xml": ("xml", 0.8),
                "application/json": ("json", 0.8),
                "text/plain": ("text", 0.7),
                "text/markdown": ("markdown", 0.8),
                "image/jpeg": ("image", 0.8),
                "image/png": ("image", 0.8),
                "image/tiff": ("image", 0.8),
                "image/bmp": ("image", 0.8),
                "image/gif": ("image", 0.8),
                "image/webp": ("image", 0.8),
                "application/zip": ("archive", 0.8),
            }

            return mime_map.get(mime_type, ("unknown", 0.0))

        except Exception:
            return "unknown", 0.0

    def _detect_by_content(self, content: str) -> Tuple[str, float]:
        """Detect file type by analyzing content"""
        if not content:
            return "unknown", 0.0

        content_lower = content.lower().strip()

        # Magic bytes and headers detection
        if content.startswith("%PDF"):
            return "pdf", 0.9
        elif content.startswith("PK"):  # ZIP-based formats (Office documents)
            return "office_document", 0.7
        elif content_lower.startswith("<!doctype html") or "<html" in content_lower:
            return "html", 0.8
        elif content_lower.startswith("<?xml"):
            return "xml", 0.8
        elif content_lower.startswith("{") and content_lower.endswith("}"):
            # JSON detection
            try:
                json.loads(content)
                return "json", 0.8
            except json.JSONDecodeError:
                pass
        elif content_lower.startswith("[") and content_lower.endswith("]"):
            # JSON array detection
            try:
                json.loads(content)
                return "json", 0.8
            except json.JSONDecodeError:
                pass
        elif self._is_csv_content(content):
            return "csv", 0.6
        elif self._is_markdown_content(content):
            return "markdown", 0.6

        # Default to text if no specific format detected
        return "text", 0.5

    def _is_csv_content(self, content: str) -> bool:
        """Check if content appears to be CSV"""
        lines = content.split("\n")[:10]  # Check first 10 lines
        if len(lines) < 2:
            return False

        # Check if most lines have commas and similar comma counts
        comma_counts = []
        for line in lines:
            if line.strip():
                comma_counts.append(line.count(","))

        if not comma_counts:
            return False

        # If most lines have the same number of commas (and > 0), likely CSV
        most_common_count = max(set(comma_counts), key=comma_counts.count)
        matching_lines = sum(1 for count in comma_counts if count == most_common_count)

        return most_common_count > 0 and matching_lines / len(comma_counts) > 0.7

    def _is_markdown_content(self, content: str) -> bool:
        """Check if content appears to be Markdown"""
        markdown_patterns = [
            r"^#+\s",  # Headers
            r"^\*\s",  # Bullet points
            r"^-\s",  # Bullet points
            r"^\d+\.\s",  # Numbered lists
            r"\*\*.*\*\*",  # Bold text
            r"\*.*\*",  # Italic text
            r"\[.*\]\(.*\)",  # Links
        ]

        lines = content.split("\n")[:20]  # Check first 20 lines
        markdown_lines = 0

        for line in lines:
            if any(re.search(pattern, line) for pattern in markdown_patterns):
                markdown_lines += 1

        return markdown_lines > len(lines) * 0.2  # If 20% of lines have markdown syntax

    def _classify_structure_with_content(
        self, file_type: str, content: str, filename: str
    ) -> str:
        """Enhanced structure classification using actual content analysis"""

        # For DOCX files, analyze the actual content structure
        if file_type == "docx" and content:
            return self._analyze_docx_structure(content)

        # For other file types, use enhanced 7-point classification
        highly_structured_types = {
            "json",
            "csv",
            "xlsx",
            "xls",
            "xml",
            "ods",
        }  # Database-like
        structured_types = {"pptx", "ppt", "odp"}  # Clear organization
        moderately_structured_types = {"html", "markdown"}  # Some structure
        semi_structured_types = {"pdf", "odt"}  # Mixed content
        lightly_structured_types = {"rtf"}  # Some formatting
        unstructured_types = {"text"}  # Plain text

        if file_type in highly_structured_types:
            return "highly_structured"
        elif file_type in structured_types:
            return "structured"
        elif file_type in moderately_structured_types:
            return "moderately_structured"
        elif file_type in semi_structured_types:
            return "semi_structured"
        elif file_type in lightly_structured_types:
            return "lightly_structured"
        elif file_type in unstructured_types:
            return "unstructured"
        elif file_type == "image":
            return "image"
        else:
            return "unknown"

    def _analyze_docx_structure(self, content: str) -> str:
        """Analyze DOCX content to determine structure type"""
        if not content or len(content.strip()) < 10:
            return "unknown"

        lines = content.split("\n")
        non_empty_lines = [line.strip() for line in lines if line.strip()]

        if len(non_empty_lines) == 0:
            return "unknown"

        # Analyze structural indicators
        structure_score = 0

        # 1. Table indicators (strong structured signal)
        table_indicators = content.count(" | ")
        if table_indicators > 20:
            structure_score += 4  # Very high table count = definitely structured
        elif table_indicators > 10:
            structure_score += 3
        elif table_indicators > 5:
            structure_score += 2
        elif table_indicators > 0:
            structure_score += 1

        # 2. Consistent formatting patterns
        # Look for repeated patterns that suggest structure
        line_patterns = {}
        for line in non_empty_lines[:20]:  # Check first 20 lines
            # Count patterns like "Field: Value"
            if ":" in line and len(line.split(":")) == 2:
                structure_score += 0.5
            # Count numbered/bulleted lists
            if re.match(r"^\d+\.\s", line) or re.match(r"^[•\-\*]\s", line):
                structure_score += 0.3

        # 3. Data-like patterns
        # Look for consistent data formats
        numeric_lines = sum(1 for line in non_empty_lines if re.search(r"\d+", line))
        if numeric_lines > len(non_empty_lines) * 0.5:
            structure_score += 1

        # 4. Headers and sections (semi-structured indicators)
        header_count = 0
        section_count = 0
        for line in non_empty_lines:
            # All caps lines (potential headers)
            if line.isupper() and len(line) > 3:
                header_count += 1
            # Lines that end with colons (potential section headers)
            elif line.endswith(":") and len(line.split()) < 6:
                section_count += 1
            # Lines that look like headings (short, title-case)
            elif len(line.split()) <= 5 and line.istitle():
                header_count += 1

        # Score based on header/section presence
        if header_count > 5:
            structure_score += 1.0  # Many headers suggests semi-structured
        elif header_count > 2:
            structure_score += 0.7
        elif header_count > 0:
            structure_score += 0.3

        if section_count > 3:
            structure_score += 0.5

        # 5. Check for long paragraphs (unstructured indicator)
        long_paragraphs = sum(1 for line in non_empty_lines if len(line) > 100)
        if long_paragraphs > len(non_empty_lines) * 0.3:
            structure_score -= 1

        # 6. Sentence complexity (unstructured indicator)
        avg_line_length = sum(len(line) for line in non_empty_lines) / len(
            non_empty_lines
        )
        if avg_line_length > 80:  # Long average line length suggests prose
            structure_score -= 0.5

        # Calculate some additional metrics for logging
        field_value_count = sum(
            1
            for line in non_empty_lines[:20]
            if ":" in line and len(line.split(":")) == 2
        )
        numeric_lines = sum(1 for line in non_empty_lines if re.search(r"\d+", line))
        long_paragraphs = sum(1 for line in non_empty_lines if len(line) > 100)
        avg_line_length = sum(len(line) for line in non_empty_lines) / len(
            non_empty_lines
        )

        logger.info(
            f"DOCX structure analysis - Score: {structure_score}, Lines: {len(non_empty_lines)}, "
            f"Tables: {table_indicators}, Field-value pairs: {field_value_count}, "
            f"Headers: {header_count}, Sections: {section_count}, "
            f"Numeric lines: {numeric_lines}, Long paragraphs: {long_paragraphs}, "
            f"Avg line length: {avg_line_length:.1f}"
        )

        # Classification based on score with 7-point scale
        if structure_score >= 4.5:
            return "highly_structured"  # 7 - Database-like, heavy tables/forms
        elif structure_score >= 3.0:
            return "structured"  # 6 - Clear data organization
        elif structure_score >= 2.0:
            return "moderately_structured"  # 5 - Some data patterns
        elif structure_score >= 1.0:
            return "semi_structured"  # 4 - Mixed content (original middle)
        elif structure_score >= 0.5:
            return "lightly_structured"  # 3 - Mostly narrative with some organization
        elif structure_score >= 0.0:
            return "unstructured"  # 2 - Prose with minimal formatting
        else:
            return "highly_unstructured"  # 1 - Pure narrative, no structure

    def _detect_scanned_content(self, file_type: str, content: str) -> bool:
        """Detect if document is scanned (image-based)"""
        # Images are always considered scanned
        if file_type == "image":
            return True

        # For DOCX files, check if content is very sparse (might be image-based)
        if file_type == "docx" and content:
            lines = content.split("\n")
            non_empty_lines = [line.strip() for line in lines if line.strip()]

            # If very few text lines but file exists, might be scanned
            if len(non_empty_lines) < 3:
                return True

            # Check if content is mostly short lines (OCR artifacts)
            short_lines = sum(1 for line in non_empty_lines if len(line.split()) < 3)
            if short_lines / len(non_empty_lines) > 0.8:
                return True

        # For PDFs, we'd need actual PDF analysis
        # For now, return False as placeholder
        if file_type == "pdf":
            return False

        return False

    def _generate_metadata(
        self,
        content: str,
        filename: str,
        file_path: str,
        content_type: str,
        file_type: str,
    ) -> Dict[str, Any]:
        """Generate metadata about the document"""
        metadata = {
            "filename": filename,
            "file_path": file_path,
            "content_type": content_type,
            "content_length": len(content) if content else 0,
            "file_type": file_type,
        }

        # Add type-specific metadata
        if file_type == "csv" and content:
            lines = content.split("\n")
            metadata["estimated_rows"] = len([line for line in lines if line.strip()])
            if lines:
                metadata["estimated_columns"] = lines[0].count(",") + 1

        elif file_type == "json" and content:
            try:
                data = json.loads(content)
                if isinstance(data, dict):
                    metadata["json_type"] = "object"
                    metadata["json_keys"] = list(data.keys())[:10]  # First 10 keys
                elif isinstance(data, list):
                    metadata["json_type"] = "array"
                    metadata["json_length"] = len(data)
            except json.JSONDecodeError:
                pass

        return metadata

    def _calculate_structure_rating(self, structure_type: str) -> str:
        """Calculate structure rating out of 7 based on classification"""
        rating_map = {
            "highly_structured": "7/7",
            "structured": "6/7",
            "moderately_structured": "5/7",
            "semi_structured": "4/7",
            "lightly_structured": "3/7",
            "unstructured": "2/7",
            "highly_unstructured": "1/7",
            "scanned": "0/7",  # Scanned documents can't be rated structurally
            "image": "0/7",
            "unknown": "0/7",
        }
        return rating_map.get(structure_type, "0/7")

    def _fallback_response(self) -> Dict[str, Any]:
        """Fallback response when detection fails"""
        return {
            "document_type": "text",
            "structure_type": "unstructured",
            "structure_rating": "2/7",
            "confidence": 0.3,
            "is_scanned": False,
            "metadata": {
                "error": "Detection failed, using fallback",
                "content_length": 0,
            },
            "success": False,
            "note": "Fallback response due to detection error",
        }
